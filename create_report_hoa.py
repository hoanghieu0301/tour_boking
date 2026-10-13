import docx
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def build_report():
    # Load original document to preserve sections, margins, styles, header/footer
    doc = docx.Document('5250050_HoangXuanHieu_CNTT_K29.1.docx')

    # Ensure section margins are correct
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.0)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)

    # Set normal style font to Times New Roman
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    # Ensure complex script font is also Times New Roman
    rPr = style_normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    # Helper function to set run font
    def set_font(run, size_pt=13, bold=False, italic=False, font_name='Times New Roman', color=None):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color
        rPr = run.element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)

    # Helper to set paragraph format
    def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=3, line_spacing=1.3):
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing

    # ==========================================
    # 1. UPDATE PARAGRAPHS 0 to 86
    # ==========================================
    # Let's map out our new paragraphs content
    new_paragraphs_data = [
        # P[0]
        {"text": "", "align": WD_ALIGN_PARAGRAPH.LEFT, "runs": []},
        # P[1] Date
        {"text": "", "align": WD_ALIGN_PARAGRAPH.RIGHT, "space_after": 6, "runs": [
            {"text": "Hà Nội, ngày 11 tháng 09 năm 2026", "size": 13, "bold": False, "italic": True}
        ]},
        # P[2]
        {"text": "", "align": WD_ALIGN_PARAGRAPH.CENTER, "runs": []},
        # P[3] Title
        {"text": "", "align": WD_ALIGN_PARAGRAPH.CENTER, "space_before": 6, "space_after": 12, "runs": [
            {"text": "ĐỀ CƯƠNG ĐỒ ÁN TỐT NGHIỆP", "size": 17, "bold": True, "italic": False}
        ]},
        # P[4]
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "runs": []},
        # P[5] Student name
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.4, "runs": [
            {"text": "Họ và tên sinh viên: ", "size": 13, "bold": True, "italic": False},
            {"text": "Nguyễn Thị Hoa", "size": 13, "bold": True, "italic": False}
        ]},
        # P[6] Student ID, class, cohort
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.4, "runs": [
            {"text": "Mã SV: ", "size": 13, "bold": True, "italic": False},
            {"text": "5250054", "size": 13, "bold": False, "italic": False},
            {"text": "                        Lớp: ", "size": 13, "bold": True, "italic": False},
            {"text": "CNTT1", "size": 13, "bold": False, "italic": False},
            {"text": "                       Khóa: ", "size": 13, "bold": True, "italic": False},
            {"text": "K29.1", "size": 13, "bold": False, "italic": False}
        ]},
        # P[7] Phone & Email
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.4, "runs": [
            {"text": "Số điện thoại\t: ", "size": 13, "bold": True, "italic": False},
            {"text": "0986543210", "size": 13, "bold": False, "italic": False},
            {"text": "                             Email: ", "size": 13, "bold": True, "italic": False},
            {"text": "hoant5250054@utc.edu.vn", "size": 13, "bold": False, "italic": False}
        ]},
        # P[8] Major
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.4, "runs": [
            {"text": "Ngành\t: ", "size": 13, "bold": True, "italic": False},
            {"text": "Công nghệ thông tin", "size": 13, "bold": False, "italic": False},
            {"text": "                Hệ: ", "size": 13, "bold": True, "italic": False},
            {"text": "Liên thông", "size": 13, "bold": False, "italic": False}
        ]},
        # P[9] Supervisor
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.4, "runs": [
            {"text": "Giảng viên (cán bộ) hướng dẫn: ", "size": 13, "bold": True, "italic": False},
            {"text": "ThS. Nguyễn Hữu Luân", "size": 13, "bold": False, "italic": False}
        ]},
        # P[10] Department
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.4, "runs": [
            {"text": "Đơn vị công tác\t: ", "size": 13, "bold": True, "italic": False},
            {"text": "Khoa Công Nghệ Thông Tin, Trường Đại học Giao Thông Vận Tải", "size": 13, "bold": False, "italic": False}
        ]},
        # P[11] Supervisor Phone & Email
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.4, "runs": [
            {"text": "Số điện thoại\t: ", "size": 13, "bold": True, "italic": False},
            {"text": "0941070588", "size": 13, "bold": False, "italic": False},
            {"text": "                               Email: ", "size": 13, "bold": True, "italic": False},
            {"text": "luan.nh@utc.edu.vn", "size": 13, "bold": False, "italic": False}
        ]},
        # P[12] Project Title
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 8, "line_spacing": 1.4, "runs": [
            {"text": "Tên đề tài: ", "size": 13, "bold": True, "italic": False},
            {"text": "Thiết kế và Xây dựng website đặt tour du lịch trực tuyến tích hợp trợ lý ảo AI", "size": 13, "bold": True, "italic": False}
        ]},
        # P[13] Section 1 Header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 8, "space_after": 4, "runs": [
            {"text": "1. Nội dung, phạm vi của đề tài", "size": 13, "bold": True, "italic": False}
        ]},
        # P[14] Content Header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 3, "runs": [
            {"text": "Nội dung", "size": 13, "bold": True, "italic": False}
        ]},
        # P[15] Content paragraph 1
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 4, "runs": [
            {"text": "Xây dựng website đặt tour du lịch trực tuyến ViViTour (Tour Booking Platform) tích hợp trợ lý ảo AI nhằm số hóa toàn diện quy trình tìm kiếm, tư vấn, đặt tour và quản lý dịch vụ lữ hành. Hệ thống giúp du khách dễ dàng tra cứu các tour du lịch trong nước và quốc tế theo điểm khởi hành, điểm đến, thời lượng, mức giá và dòng tour (Cao cấp, Tiêu chuẩn, Tiết kiệm); theo dõi chi tiết lịch trình theo từng ngày (itinerary), dịch vụ bao gồm/không bao gồm, chính sách hủy tour và số chỗ còn trống theo thời gian thực; loại bỏ hoàn toàn sự chậm trễ của phương thức liên hệ và tư vấn thủ công truyền thống.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[16] Content paragraph 2
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 4, "runs": [
            {"text": "Hệ thống cho phép khách hàng đăng ký, đăng nhập tài khoản (hỗ trợ xác thực bảo mật chuẩn JWT và đăng nhập nhanh bằng tài khoản Google OAuth 2.0), tìm kiếm và lọc tour đa tiêu chí, thực hiện đặt tour trực tuyến với thông tin hành khách phân loại chi tiết (người lớn, trẻ em, em bé và tùy chọn phụ thu phòng đơn), áp dụng mã giảm giá (voucher/promo code) và thanh toán trực tuyến an toàn qua cổng VNPay. Khách hàng nhận email xác nhận đặt tour kèm hóa đơn điện tử tự động qua dịch vụ Nodemailer, theo dõi lịch sử đơn đặt, gửi yêu cầu hủy tour, lưu danh sách tour yêu thích (wishlist) và tham gia đánh giá, bình luận chất lượng tour sau chuyến đi. Đặc biệt, khách hàng được hỗ trợ tư vấn 24/7 bởi Trợ lý ảo AI thông minh (tích hợp OpenAI / Google Gemini API) có khả năng gợi ý lịch trình và giải đáp thắc mắc phù hợp với nhu cầu và ngân sách.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[17] Content paragraph 3
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 4, "runs": [
            {"text": "Phía quản trị viên (Admin) được cung cấp hệ thống quản trị chuyên sâu (Admin Dashboard): bảng điều khiển thống kê trực quan doanh thu, số lượt đặt tour, người dùng mới và tỷ lệ lấp đầy lịch khởi hành; quản lý danh mục tour và cấu hình lịch khởi hành linh hoạt (TourSchedule: ngày khởi hành, số chỗ, giá vé theo ngày); quản lý chuỗi cung ứng dịch vụ đối tác (Supplier: phương tiện vận chuyển, khách sạn lưu trú, nhà hàng ẩm thực, hướng dẫn viên); quản lý đơn đặt tour (Booking) và xét duyệt yêu cầu hủy tour/hoàn tiền; quản lý danh mục địa điểm (tỉnh thành, khu vực); cấu hình chiến dịch khuyến mãi (Promo Code) và kiểm duyệt các đánh giá của khách hàng.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[18] Content paragraph 4 (Empty or brief summary)
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 4, "runs": [
            {"text": "Hệ thống được thiết kế theo kiến trúc hiện đại, phân tách rõ ràng giữa phân hệ Backend API (NestJS, Prisma ORM, MySQL) và hai phân hệ Frontend độc lập (User Web & Admin Web xây dựng trên Next.js 15, React 19, TailwindCSS), đảm bảo tính module hóa cao, dễ dàng bảo trì, mở rộng và tối ưu hóa hiệu năng trải nghiệm người dùng.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[19] Scope header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 6, "space_after": 3, "runs": [
            {"text": "Phạm vi đề tài", "size": 13, "bold": True, "italic": False}
        ]},
        # P[20] Tech scope header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 3, "space_after": 2, "runs": [
            {"text": "Công nghệ và nền tảng sử dụng:", "size": 13, "bold": True, "italic": False}
        ]},
        # P[21] Backend tech
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Ngôn ngữ lập trình & Framework: ", "size": 13, "bold": True, "italic": False},
            {"text": "Sử dụng TypeScript trên nền tảng Node.js kết hợp framework NestJS 11 để xây dựng hệ thống Backend kiến trúc phân tầng module hóa (Controller, Service, Repository, Module, Guard, Interceptor), xử lý các nghiệp vụ đặt tour du lịch, tích hợp AI và giao dịch thanh toán trực tuyến an toàn.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[22] Database tech
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Cơ sở dữ liệu & ORM: ", "size": 13, "bold": True, "italic": False},
            {"text": "Sử dụng hệ quản trị CSDL quan hệ MySQL kết hợp Prisma ORM v7 để định nghĩa Schema, quản lý migration và thao tác dữ liệu an toàn, đảm bảo tính toàn vẹn (ACID) cho toàn bộ thông tin người dùng, tour, lịch trình, hành khách, đơn đặt, thanh toán và đối tác cung ứng.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[23] Frontend tech
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Giao diện người dùng: ", "size": 13, "bold": True, "italic": False},
            {"text": "Sử dụng framework Next.js 15 (App Router, Turbopack) và React 19 kết hợp TailwindCSS v4, TypeScript, Lucide React và Recharts để phát triển giao diện người dùng (Client Web) và quản trị viên (Admin Web) responsive, hiện đại, tối ưu trải nghiệm trên mọi thiết bị.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[24] Third-party tech
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Tích hợp dịch vụ bên thứ ba & Trí tuệ nhân tạo: ", "size": 13, "bold": True, "italic": False},
            {"text": "Tích hợp cổng thanh toán trực tuyến VNPay (Sandbox) xử lý giao dịch an toàn với mã hóa HMAC-SHA512; tích hợp Google Gemini API & OpenAI API phục vụ Trợ lý ảo AI tư vấn tour; tích hợp Cloudinary API & AWS S3 lưu trữ đa phương tiện; sử dụng Nodemailer (SMTP) gửi mã OTP và vé tour điện tử; tích hợp Google OAuth 2.0 đăng nhập tiện lợi.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[25] Dev tools
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Công cụ phát triển: ", "size": 13, "bold": True, "italic": False},
            {"text": "Sử dụng Visual Studio Code / Antigravity IDE làm môi trường phát triển chính. Quản lý thư viện bằng NPM. Sử dụng Postman để thiết kế và kiểm thử API. Quản lý mã nguồn và phiên bản bằng Git/GitHub. Sử dụng Prisma Studio để quản trị dữ liệu trực quan.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[26] Deployment
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Triển khai hệ thống: ", "size": 13, "bold": True, "italic": False},
            {"text": "Chạy hệ thống trên môi trường máy chủ cục bộ (Localhost: User Web port 3000, Admin Web port 3001, Backend API port 8000/8080, MySQL port 3306) trong quá trình phát triển và kiểm thử; có khả năng đóng gói Docker và triển khai lên máy chủ đám mây (Cloud Server/VPS).", "size": 13, "bold": False, "italic": False}
        ]},
        # P[27] Data scope header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 2, "runs": [
            {"text": "Quy mô dữ liệu và đối tượng người dùng:", "size": 13, "bold": True, "italic": False}
        ]},
        # P[28] Data scope details
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Dữ liệu bao gồm: thông tin tài khoản người dùng, danh mục địa điểm du lịch (vùng miền, tỉnh thành, điểm đến), danh mục tour du lịch, lịch khởi hành theo ngày (TourSchedule), lịch trình chi tiết (TourItinerary), thông tin đối tác cung ứng dịch vụ (Supplier: vận chuyển, khách sạn, nhà hàng, hướng dẫn viên), đơn đặt tour (Booking), danh sách hành khách chi tiết, hóa đơn thanh toán VNPay, mã khuyến mãi (PromoCode), đánh giá nhận xét và lịch sử hội thoại tư vấn AI.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[29] (Merge or brief)
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 2, "space_after": 2, "runs": [
            {"text": "Hệ thống phục vụ hai nhóm đối tượng chính:", "size": 13, "bold": True, "italic": False}
        ]},
        # P[30] User groups
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "1. Khách hàng (Du khách): ", "size": 13, "bold": True, "italic": False},
            {"text": "Tìm kiếm, xem thông tin tour, nhận tư vấn trực tiếp từ trợ lý ảo AI, đặt tour, điền thông tin hành khách, thanh toán VNPay, quản lý lịch sử đặt và viết đánh giá.\n", "size": 13, "bold": False, "italic": False},
            {"text": "2. Quản trị viên (Admin): ", "size": 13, "bold": True, "italic": False},
            {"text": "Quản lý toàn diện hệ thống du lịch, danh mục tour, lịch khởi hành, quản lý chuỗi cung ứng dịch vụ, kiểm soát đơn đặt, duyệt hủy tour, cấu hình khuyến mãi và theo dõi báo cáo doanh thu.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[31] Processing scope header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 2, "runs": [
            {"text": "Phạm vi xử lý:", "size": 13, "bold": True, "italic": False}
        ]},
        # P[32] Processing scope details
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Hệ thống tập trung vào các quy trình nghiệp vụ cốt lõi: tìm kiếm và lọc tour đa tiêu chí theo thời gian thực; hiển thị lịch trình chi tiết và tình trạng chỗ còn trống; tư vấn du lịch tự động bằng Trợ lý ảo AI; quy trình đặt tour và nhập danh sách hành khách; thanh toán trực tuyến bảo mật qua cổng VNPay; tự động gửi email xác nhận đặt tour và vé điện tử; tiếp nhận và xử lý yêu cầu hủy tour; quản lý danh sách yêu thích và đánh giá sau tour; cùng bảng điều khiển quản trị Admin toàn diện.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[33] Limitations header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 2, "runs": [
            {"text": "Giới hạn của đề tài:", "size": 13, "bold": True, "italic": False}
        ]},
        # P[34] Limitations 1
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Đề tài tập trung vào bài toán phát triển nền tảng ứng dụng web (Web Application) quản lý và đặt tour du lịch trực tuyến. Chưa phát triển ứng dụng di động độc lập (Native Mobile App) trên nền tảng iOS/Android (tuy nhiên giao diện web được thiết kế hoàn toàn tương thích và tối ưu hiển thị responsive trên các thiết bị di động).", "size": 13, "bold": False, "italic": False}
        ]},
        # P[35] Limitations 2
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Cổng thanh toán VNPay được thử nghiệm và mô phỏng trên môi trường Sandbox theo chuẩn kết nối API của đơn vị cung cấp thanh toán. Hệ thống email tự động sử dụng giao thức SMTP tiêu chuẩn của Google. Trợ lý ảo AI sử dụng API của OpenAI / Google Gemini kết hợp cơ sở dữ liệu tour có sẵn.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[36] Limitations 3
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "Hệ thống cung cấp các báo cáo thống kê trực quan về doanh thu, số lượng đơn đặt tour theo ngày/tháng/năm và các tour bán chạy; chưa tích hợp các thuật toán học sâu dự báo biến động giá tour nâng cao.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[37] Extensibility header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 2, "runs": [
            {"text": "Khả năng mở rộng:", "size": 13, "bold": True, "italic": False}
        ]},
        # P[38] Extensibility details
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 4, "runs": [
            {"text": "Trong tương lai, hệ thống có thể phát triển thêm ứng dụng di động đa nền tảng (Mobile App trên nền Flutter / React Native), tích hợp thêm các cổng thanh toán quốc tế (PayPal, Stripe) và ví điện tử (MoMo, ZaloPay), tích hợp hệ thống bản đồ số dẫn đường tương tác thực tế (Interactive GPS Map), mô hình AI cá nhân hóa hành trình du lịch (Smart Travel Planner) và mở rộng kết nối API trực tiếp với các hệ thống khách sạn, hãng hàng không toàn cầu.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[39] Technologies header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 8, "space_after": 4, "runs": [
            {"text": "Công nghệ, công cụ và ngôn ngữ lập trình", "size": 13, "bold": True, "italic": False}
        ]},
        # P[40] Frontend header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 3, "space_after": 2, "runs": [
            {"text": "a) Front-end", "size": 13, "bold": True, "italic": False}
        ]},
        # P[41] Frontend 1
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Next.js 15 & React 19: ", "size": 13, "bold": True, "italic": False},
            {"text": "Xây dựng ứng dụng web hiện đại theo cơ chế App Router, hỗ trợ Server-Side Rendering (SSR) và Client-Side Rendering (CSR) tối ưu tốc độ tải trang và SEO.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[42] Frontend 2
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " TailwindCSS v4: ", "size": 13, "bold": True, "italic": False},
            {"text": "Framework CSS tiện ích giúp xây dựng giao diện hiện đại, responsive, thẩm mỹ cao và đồng bộ nhận diện thương hiệu trên mọi thiết bị.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[43] Frontend 3
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " TypeScript: ", "size": 13, "bold": True, "italic": False},
            {"text": "Ngôn ngữ lập trình định kiểu tĩnh, nâng cao độ tin cậy, giảm thiểu lỗi runtime và tối ưu hóa năng suất phát triển giao diện người dùng.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[44] Frontend 4
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": " Lucide React & Recharts: ", "size": 13, "bold": True, "italic": False},
            {"text": "Bộ biểu tượng vector hiện đại kết hợp thư viện biểu đồ trực quan hóa dữ liệu thống kê doanh thu và đơn đặt tour trên Admin Dashboard.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[45] Backend header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 2, "runs": [
            {"text": "b) Back-end", "size": 13, "bold": True, "italic": False}
        ]},
        # P[46] Backend 1
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " TypeScript / Node.js: ", "size": 13, "bold": True, "italic": False},
            {"text": "Môi trường thực thi JavaScript phía máy chủ kết hợp ngôn ngữ TypeScript định kiểu chặt chẽ, tối ưu hiệu năng xử lý bất đồng bộ.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[47] Backend 2
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " NestJS Framework (v11): ", "size": 13, "bold": True, "italic": False},
            {"text": "Framework xây dựng hệ thống Backend kiến trúc phân tầng module hóa chuẩn doanh nghiệp, hỗ trợ Dependency Injection mạnh mẽ và bảo trì dễ dàng.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[48] Backend 3
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Prisma ORM (v7): ", "size": 13, "bold": True, "italic": False},
            {"text": "Công cụ quản lý mô hình dữ liệu Type-safe thế hệ mới, tự động sinh mã tương tác CSDL, quản lý migrations và tối ưu truy vấn.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[49] Backend 4
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Passport & JWT (JSON Web Token): ", "size": 13, "bold": True, "italic": False},
            {"text": "Xác thực và phân quyền người dùng an toàn (Access Token, Refresh Token, Blacklist Token, Role-based Access Control cho Admin & User).", "size": 13, "bold": False, "italic": False}
        ]},
        # P[50] Backend 5
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Cổng thanh toán VNPay API: ", "size": 13, "bold": True, "italic": False},
            {"text": "Tích hợp cổng thanh toán trực tuyến VNPay Sandbox với chuẩn bảo mật SHA-512, xử lý tự động giao dịch qua Return URL và IPN webhook.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[51] Backend 6
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Trí tuệ nhân tạo (AI / LLM): ", "size": 13, "bold": True, "italic": False},
            {"text": "Tích hợp OpenAI API và Google Gemini API (@google/genai) kết hợp lưu trữ vector embedding (TourEmbedding) phục vụ tư vấn tour thông minh.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[52] Backend 7
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": " Nodemailer, Cloudinary & AWS S3: ", "size": 13, "bold": True, "italic": False},
            {"text": "Tích hợp Nodemailer gửi email hóa đơn và vé tour tự động; sử dụng Cloudinary và AWS S3 lưu trữ hình ảnh/video tour chất lượng cao.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[53] Database header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 2, "runs": [
            {"text": "c) Cơ sở dữ liệu (Database)", "size": 13, "bold": True, "italic": False}
        ]},
        # P[54] Database content
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": "MySQL Server: ", "size": 13, "bold": True, "italic": False},
            {"text": "Hệ quản trị CSDL quan hệ lưu trữ toàn bộ dữ liệu quan hệ của hệ thống (User, Tour, TourSchedule, TourItinerary, Booking, Payment, PromoCode, Supplier, Review, ChatSession...), đảm bảo tính toàn vẹn (ACID) và hiệu năng truy vấn cao.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[55] Tools header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 2, "runs": [
            {"text": "d) Công cụ phát triển", "size": 13, "bold": True, "italic": False}
        ]},
        # P[56] Tool 1
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Visual Studio Code / Antigravity IDE: ", "size": 13, "bold": True, "italic": False},
            {"text": "Môi trường phát triển tích hợp (IDE) chính cho ứng dụng TypeScript, Next.js và NestJS.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[57] Tool 2
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " NPM / NPX: ", "size": 13, "bold": True, "italic": False},
            {"text": "Công cụ quản lý gói thư viện phụ thuộc và thực thi các tập lệnh dự án.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[58] Tool 3
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Git & GitHub: ", "size": 13, "bold": True, "italic": False},
            {"text": "Hệ thống quản lý mã nguồn phân tán và kiểm soát phiên bản dự án.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[59] Tool 4
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 3, "runs": [
            {"text": " Postman & Prisma Studio: ", "size": 13, "bold": True, "italic": False},
            {"text": "Kiểm thử hệ thống RESTful API và quản trị trực quan cơ sở dữ liệu trên giao diện web.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[60] Deployment env header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 2, "runs": [
            {"text": "e) Môi trường triển khai", "size": 13, "bold": True, "italic": False}
        ]},
        # P[61] Env 1
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": "Node.js Runtime Environment (v20+ / v22+): ", "size": 13, "bold": True, "italic": False},
            {"text": "Môi trường thực thi ứng dụng JavaScript/TypeScript.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[62] Env 2
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": "Next.js Web Server & NestJS API Server: ", "size": 13, "bold": True, "italic": False},
            {"text": "Máy chủ phục vụ giao diện người dùng và xử lý nghiệp vụ API.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[63] Env 3
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 4, "runs": [
            {"text": "MySQL Server: ", "size": 13, "bold": True, "italic": False},
            {"text": "Hệ thống quản lý cơ sở dữ liệu quan hệ.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[64] Expected results header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 8, "space_after": 4, "runs": [
            {"text": "Các kết quả chính dự kiến đạt được", "size": 13, "bold": True, "italic": False}
        ]},
        # P[65] Customer section header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 4, "space_after": 3, "runs": [
            {"text": "* Đối với khách hàng (Du khách):", "size": 13, "bold": True, "italic": False}
        ]},
        # P[66] Customer feature 1
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Đăng ký / Đăng nhập: ", "size": 13, "bold": True, "italic": False},
            {"text": "Đăng ký tài khoản, đăng nhập an toàn với mật khẩu băm Bcrypt; hỗ trợ đăng nhập một chạm qua Google OAuth 2.0; tính năng quên mật khẩu và xác thực qua mã OTP email.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[67] Customer feature 2
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Tra cứu & Lọc tour đa tiêu chí: ", "size": 13, "bold": True, "italic": False},
            {"text": "Tìm kiếm tour theo từ khóa, điểm khởi hành, điểm đến; lọc theo mức giá, thời lượng ngày, dòng tour (Cao cấp, Tiêu chuẩn, Tiết kiệm) và phương tiện di chuyển.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[68] Customer feature 3
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Xem chi tiết tour & Lịch trình: ", "size": 13, "bold": True, "italic": False},
            {"text": "Xem thư viện ảnh/video chất lượng cao, lịch trình chi tiết theo từng ngày (itinerary), dịch vụ bao gồm/không bao gồm, thông tin khách sạn, bữa ăn và chính sách hủy tour.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[69] Customer feature 4
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Đặt tour trực tuyến & Nhập thông tin hành khách: ", "size": 13, "bold": True, "italic": False},
            {"text": "Chọn ngày khởi hành còn trống, nhập thông tin liên hệ và danh sách hành khách phân loại chi tiết (người lớn, trẻ em, em bé), tùy chọn phụ thu phòng đơn với giá tính tự động.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[70] Customer feature 5
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Thanh toán trực tuyến VNPay: ", "size": 13, "bold": True, "italic": False},
            {"text": "Tích hợp cổng VNPay xử lý thanh toán trực tuyến nhanh chóng, bảo mật tuyệt đối và xác nhận thanh toán tức thời.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[71] Customer feature 6
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Quản lý mã giảm giá (Promo Code): ", "size": 13, "bold": True, "italic": False},
            {"text": "Áp dụng các mã voucher giảm giá (theo số tiền cố định hoặc phần trăm) trực tiếp vào đơn đặt tour để nhận ưu đãi.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[72] Customer feature 7
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Email xác nhận & Vé điện tử: ", "size": 13, "bold": True, "italic": False},
            {"text": "Hệ thống tự động gửi email hóa đơn xác nhận đặt tour và thông tin vé điện tử chi tiết ngay sau khi thanh toán thành công.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[73] Customer feature 8
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Trợ lý ảo AI Chatbot (24/7): ", "size": 13, "bold": True, "italic": False},
            {"text": "Tư vấn viên ảo hỗ trợ trò chuyện, giải đáp thắc mắc và gợi ý các tour du lịch phù hợp nhất theo sở thích và ngân sách của khách hàng.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[74] Customer feature 9
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 4, "runs": [
            {"text": " Quản lý lịch sử đặt, Hủy tour & Đánh giá: ", "size": 13, "bold": True, "italic": False},
            {"text": "Theo dõi trạng thái đơn đặt tour, gửi yêu cầu hủy tour, quản lý danh sách tour yêu thích (wishlist), cập nhật thông tin cá nhân và gửi đánh giá (rating 1-5 sao kèm nhận xét) sau chuyến đi.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[75] Admin section header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 6, "space_after": 3, "runs": [
            {"text": "* Đối với quản trị viên (Admin):", "size": 13, "bold": True, "italic": False}
        ]},
        # P[76] Admin feature 1
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Bảng điều khiển (Dashboard) thống kê: ", "size": 13, "bold": True, "italic": False},
            {"text": "Biểu đồ trực quan thống kê tổng doanh thu, số lượt đặt tour, khách hàng mới, tỷ lệ lấp đầy lịch khởi hành và các tour du lịch được yêu thích nhất.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[77] Admin feature 2
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Quản lý Tour & Lịch trình: ", "size": 13, "bold": True, "italic": False},
            {"text": "Thêm mới, chỉnh sửa, xóa và cấu hình thông tin tour (giá gốc, thời lượng, phương tiện, dòng tour, dịch vụ bao gồm/không bao gồm, lộ trình chi tiết theo từng ngày, khách sạn và bữa ăn).", "size": 13, "bold": False, "italic": False}
        ]},
        # P[78] Admin feature 3
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Quản lý Lịch khởi hành (TourSchedule): ", "size": 13, "bold": True, "italic": False},
            {"text": "Thiết lập các đợt khởi hành cho từng tour, số lượng chỗ ngồi mở bán, giá vé điều chỉnh theo ngày và trạng thái mở/đóng lịch khởi hành.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[79] Admin feature 4
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Quản lý Nhà cung cấp (Supplier Management): ", "size": 13, "bold": True, "italic": False},
            {"text": "Quản lý danh sách đối tác cung cấp dịch vụ (đơn vị vận chuyển, khách sạn lưu trú, nhà hàng ẩm thực, hướng dẫn viên) và phân bổ vào từng chặng tour.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[80] Admin feature 5
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Quản lý Đơn đặt tour (Booking Management): ", "size": 13, "bold": True, "italic": False},
            {"text": "Tra cứu, lọc đơn theo trạng thái (chờ thanh toán, đã xác nhận, hoàn thành, hủy), kiểm tra thông tin danh sách hành khách, tiếp nhận và xét duyệt/từ chối các yêu cầu hủy tour.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[81] Admin feature 6
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Quản lý Mã giảm giá (Promo Code): ", "size": 13, "bold": True, "italic": False},
            {"text": "Tạo các chiến dịch khuyến mãi, mã voucher giảm giá (theo số tiền hoặc %), quy định số tiền giảm tối đa, thời hạn hiệu lực và tour áp dụng.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[82] Admin feature 7
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "line_spacing": 1.35, "space_after": 2, "runs": [
            {"text": " Quản lý Người dùng, Đánh giá & Danh mục địa điểm: ", "size": 13, "bold": True, "italic": False},
            {"text": "Quản lý danh sách tài khoản, khóa/mở khóa người dùng; kiểm duyệt đánh giá và phản hồi của khách hàng; quản lý danh mục địa điểm, tỉnh thành và khu vực du lịch.", "size": 13, "bold": False, "italic": False}
        ]},
        # P[83] Schedule section header
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "space_before": 8, "space_after": 6, "runs": [
            {"text": "Kế hoạch thực hiện đề tài", "size": 13, "bold": True, "italic": False}
        ]},
        # P[84] Empty spacing paragraph
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "runs": []},
        # P[85] Empty spacing paragraph
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "runs": []},
        # P[86] Empty spacing paragraph
        {"text": "", "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "runs": []}
    ]

    # Apply to document paragraphs
    for idx, p_data in enumerate(new_paragraphs_data):
        if idx < len(doc.paragraphs):
            p = doc.paragraphs[idx]
            p.text = ""  # Clear existing runs
            format_p(
                p,
                align=p_data.get("align", WD_ALIGN_PARAGRAPH.JUSTIFY),
                space_before=p_data.get("space_before", 0),
                space_after=p_data.get("space_after", 3),
                line_spacing=p_data.get("line_spacing", 1.3)
            )
            for r_info in p_data.get("runs", []):
                r = p.add_run(r_info["text"])
                set_font(
                    r,
                    size_pt=r_info.get("size", 13),
                    bold=r_info.get("bold", False),
                    italic=r_info.get("italic", False)
                )

    # ==========================================
    # 2. UPDATE TABLE 1: SCHEDULE TABLE
    # ==========================================
    t1 = doc.tables[1]
    
    # Table 1 content for Tour Booking Project
    schedule_data = [
        ("STT", "Nội dung công việc", "Thời gian dự kiến", "Ghi chú"),
        ("1", "Khảo sát đề tài, tìm hiểu thực trạng thị trường đặt tour du lịch trực tuyến, xác định mục tiêu và phạm vi hệ thống ViViTour", "10/09/2026 – 15/09/2026", ""),
        ("2", "Thu thập và phân tích yêu cầu nghiệp vụ của Khách hàng (du khách) và Quản trị viên hệ thống (Admin)", "19/09/2026 – 27/09/2026", ""),
        ("3", "Phân tích yêu cầu chức năng hệ thống: tìm kiếm tour, đặt tour & hành khách, thanh toán VNPay, trợ lý ảo AI Chatbot, quản lý nhà cung cấp", "28/09/2026 – 04/10/2026", ""),
        ("4", "Xây dựng tài liệu thiết kế hệ thống: sơ đồ ca sử dụng (Use Case), sơ đồ hoạt động (Activity), sơ đồ tuần tự (Sequence Diagram) cho các quy trình nghiệp vụ chính", "05/10/2026 – 10/10/2026", ""),
        ("5", "Thiết kế mô hình cơ sở dữ liệu quan hệ (ERD), xác định các bảng (User, Tour, TourSchedule, TourItinerary, Booking, Payment, PromoCode, Supplier...)", "11/10/2026 – 16/10/2026", ""),
        ("6", "Thiết kế kiến trúc hệ thống Backend NestJS 11, cấu hình Prisma ORM, cấu hình bảo mật Passport JWT & Google OAuth 2.0 theo chuẩn phân tầng module", "17/10/2026 – 23/10/2026", ""),
        ("7", "Khởi tạo và thiết kế giao diện Frontend Client và Admin bằng Next.js 15, React 19, TailwindCSS v4, Lucide React và thư viện Recharts", "24/10/2026 – 29/10/2026", ""),
        ("8", "Cài đặt CSDL MySQL, chạy migration với Prisma, thiết lập ràng buộc khóa ngoại, chỉ mục (Index) và viết script khởi tạo dữ liệu mẫu (Seeding)", "30/10/2026 – 05/11/2026", ""),
        ("9", "Phát triển phân hệ Khách hàng: tra cứu, tìm kiếm và lọc tour đa tiêu chí, xem chi tiết lịch trình theo ngày, quản lý tour yêu thích và hồ sơ cá nhân", "06/11/2026 – 13/11/2026", ""),
        ("10", "Phát triển phân hệ Đặt tour & Tích hợp cổng thanh toán VNPay: chọn ngày khởi hành, nhập danh sách hành khách, tạo URL thanh toán Sandbox, xử lý Return URL & IPN", "14/11/2026 – 22/11/2026", ""),
        ("11", "Tích hợp Trợ lý ảo AI Chatbot (kết nối OpenAI / Gemini API, xử lý ngữ cảnh dữ liệu tour, gợi ý tour thông minh) và dịch vụ gửi Email tự động qua Nodemailer", "23/11/2026 – 01/12/2026", ""),
        ("12", "Phát triển phân hệ Quản trị Admin: Dashboard thống kê, quản lý Tour, lịch khởi hành (TourSchedule), quản lý nhà cung cấp (Supplier), đơn đặt và mã giảm giá", "02/12/2026 – 05/12/2026", ""),
        ("13", "Phát triển chức năng Đánh giá / Nhận xét (Review), tiếp nhận và xử lý yêu cầu hủy tour, hoàn thiện các tương tác người dùng thời gian thực", "06/12/2026 – 11/12/2026", ""),
        ("14", "Kiểm thử toàn diện chức năng, kiểm thử tích hợp VNPay & Email & AI Chatbot, sửa lỗi hệ thống, tối ưu hiệu năng và chuẩn bị dữ liệu demo", "12/12/2026 – 14/12/2026", ""),
        ("15", "Hoàn thiện báo cáo đồ án tốt nghiệp, slide thuyết trình và chuẩn bị bảo vệ đề tài trước hội đồng", "15/12/2026 – 25/12/2026", "")
    ]

    for row_idx, row_data in enumerate(schedule_data):
        row = t1.rows[row_idx]
        is_header = (row_idx == 0)
        for col_idx, text in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.text = "" # clear
            p = cell.paragraphs[0]
            format_p(
                p,
                align=WD_ALIGN_PARAGRAPH.CENTER if (is_header or col_idx in [0, 2, 3]) else WD_ALIGN_PARAGRAPH.JUSTIFY,
                space_before=2,
                space_after=2,
                line_spacing=1.15
            )
            r = p.add_run(text)
            set_font(
                r,
                size_pt=13 if is_header else 11,
                bold=is_header,
                italic=False
            )

    # ==========================================
    # 3. UPDATE TABLE 2: SIGNATURES TABLE
    # ==========================================
    t2 = doc.tables[2]
    # Update Cell (0, 3) with student name Nguyen Thi Hoa
    c3 = t2.rows[0].cells[3]
    c3.text = "" # clear
    
    # Paragraph 1: Title
    p_title = c3.paragraphs[0]
    format_p(p_title, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, line_spacing=1.15)
    r_t1 = p_title.add_run("Sinh viên thực hiện")
    set_font(r_t1, size_pt=12, bold=True)
    
    # Paragraph 2: Subtitle
    p_sub = c3.add_paragraph()
    format_p(p_sub, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, line_spacing=1.15)
    r_sub = p_sub.add_run("(Ký và ghi rõ họ tên)")
    set_font(r_sub, size_pt=11, bold=False, italic=True)

    # Spacing paragraphs
    for _ in range(3):
        p_sp = c3.add_paragraph()
        format_p(p_sp, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, line_spacing=1.15)

    # Student name
    p_name = c3.add_paragraph()
    format_p(p_name, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, line_spacing=1.15)
    r_name = p_name.add_run("Nguyễn Thị Hoa")
    set_font(r_name, size_pt=12, bold=True)

    output_filename = "5250054_NguyenThiHoa_CNTT_K29.1.docx"
    doc.save(output_filename)
    print(f"Successfully generated {output_filename}")

if __name__ == "__main__":
    build_report()
